"""stories-in-wx 回归测试（轻量、自包含、秒级）。

全部使用**合成的临时数据**，不读取任何真实聊天库，因此运行很快、
不会让磁盘/CPU 长时间满载。

覆盖本轮修复的逻辑与性能缺陷：
  1. exporter._safe_name() 恒返回 "_"  → 导出名里的联系人名称丢失
  2. 媒体扩展名与 media 引用不一致（PNG/GIF 指向不存在的 .jpg）
  3. 媒体解密未传 chat/ts → attach/Bubble/Thumb 三级来源全部失效
  4. HTML 导出 session 键不匹配 → KeyError: 'displayName'（该格式不可用）
  5. session["_avatar_map"] 泄漏进导出 JSON
  6. api_chat 的 CDATA 替换串是控制字符 0x01 而非捕获组
  7. messages 的 limit 未做下限校验（LIMIT -2 等同无限制）
  8. 分片索引：消除「每个会话扫全部 *.db」的性能瓶颈
  9. 聊天页只扫 message_*.db → 漏掉 biz_message_*.db 里的会话
 10. cli --json 未生效
 11. decrypt_database 非原子写（失败会破坏已有明文库）

运行：
    python -m unittest discover -s tests -v
"""
import hashlib
import hmac as hmac_mod
import io
import json
import os
import shutil
import sqlite3
import struct
import sys
import tempfile
import time
import unittest
from unittest import mock
from contextlib import redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

# 宿主回归测试必须与"用户装了什么插件"解耦：插件可以（合法地）改写会话列表、
# 消息形状与导出产物，若参与本套件会让宿主行为断言变得不确定。
# 故在导入 siwx 之前强制零插件模式；插件自身的测试见 tests/test_plugins.py。
os.environ["SIWX_NO_PLUGINS"] = "1"

from siwx import api_chat, exporter, paths
from siwx.exporter import _safe_name


# ── 合成数据构造 ────────────────────────────────────────────────

def _msg_table(chat):
    return "Msg_" + hashlib.md5(chat.encode()).hexdigest()


def make_shard(path: Path, chat, texts, start_ts=1_700_000_000,
               include_images=False, origin=0):
    """建一个含 Msg_ 表 + Name2Id 的分片库。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    t = _msg_table(chat)
    conn.execute(f"""CREATE TABLE [{t}] (
        local_id INTEGER PRIMARY KEY, server_id INTEGER, local_type INTEGER,
        create_time INTEGER, origin_source INTEGER, real_sender_id INTEGER,
        message_content BLOB, packed_info_data BLOB)""")
    conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
    conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
    for i, txt in enumerate(texts):
        ts = start_ts + i
        if include_images and i % 5 == 0:
            content = ('<msg><img aeskey="x" md5="'
                       + hashlib.md5(f"img{i}".encode()).hexdigest()
                       + '"/></msg>').encode("utf-8")
            ltype = 3
        else:
            content = txt.encode("utf-8")
            ltype = 1
        conn.execute(f"INSERT INTO [{t}] VALUES (?,?,?,?,?,?,?,?)",
                     (i + 1, 1000 + i, ltype, ts, origin, 1, content, None))
    conn.commit()
    conn.close()
    return path


def make_empty_shard(path: Path):
    """建一个不含 Msg_ 表的分片（模拟 media_*/fts/resource 等）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE unrelated (x INTEGER)")
    conn.commit()
    conn.close()
    return path


def make_account(root: Path, account="wxid_test", chat="wxid_friend",
                 n_texts=12, include_images=False):
    """构造一个最小可用的解密产物目录。"""
    acc = root / "output" / account
    msg_dir = acc / "message"
    # 两个分片：一个含目标会话，一个不含（验证索引会跳过它）
    make_shard(msg_dir / "message_0.db", chat,
               [f"第 {i} 条消息" for i in range(n_texts)],
               include_images=include_images)
    make_empty_shard(msg_dir / "media_0.db")
    make_empty_shard(msg_dir / "message_fts.db")

    (acc / "contact").mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(acc / "contact" / "contact.db")
    c.execute("CREATE TABLE contact (username TEXT, remark TEXT, "
              "nick_name TEXT, alias TEXT)")
    c.execute("INSERT INTO contact VALUES (?,?,?,?)", (chat, "测试好友", "", ""))
    c.commit()
    c.close()

    (acc / "session").mkdir(parents=True, exist_ok=True)
    s = sqlite3.connect(acc / "session" / "session.db")
    s.execute("CREATE TABLE SessionTable (username TEXT, summary TEXT, "
              "sort_timestamp INTEGER)")
    s.execute("INSERT INTO SessionTable VALUES (?,?,?)", (chat, "预览", 1_700_000_010))
    s.commit()
    s.close()
    return acc, account, chat


class TempRootCase(unittest.TestCase):
    """把 SIWX_ROOT 指向临时目录，避免污染真实 output/exports。

    同时强制"零插件"状态：本套件断言的是宿主默认行为，若 `unittest discover`
    先跑了 tests/test_plugins.py，模块级 registry 单例里会残留合成插件的 hook
    （会话过滤器会剔掉公众号、渲染器会改写 kind），必须在此清空。
    """

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="siwx_test_"))
        self._old = os.environ.get("SIWX_ROOT")
        os.environ["SIWX_ROOT"] = str(self.tmp)
        os.environ["SIWX_NO_PLUGINS"] = "1"
        self._clear_plugins()
        api_chat._SHARD_INDEX.clear()
        api_chat._CONTACT_CACHE.clear()
        api_chat._SESSION_CACHE.clear()
        paths._PATH_CACHE.clear()

    def tearDown(self):
        if self._old is None:
            os.environ.pop("SIWX_ROOT", None)
        else:
            os.environ["SIWX_ROOT"] = self._old
        os.environ["SIWX_NO_PLUGINS"] = "1"
        self._clear_plugins()
        api_chat._SHARD_INDEX.clear()
        api_chat._CONTACT_CACHE.clear()
        api_chat._SESSION_CACHE.clear()
        paths._PATH_CACHE.clear()
        shutil.rmtree(self.tmp, ignore_errors=True)

    @staticmethod
    def _clear_plugins() -> None:
        """清空模块级 registry（插件测试可能留下合成 hook）。"""
        try:
            from siwx.plugins.registry import registry
        except Exception:
            return
        for ns in vars(registry).values():
            if hasattr(ns, "items") and isinstance(ns.items, list):
                ns.items = []
        registry.metas = {}
        registry.report = None
        registry._loaded = True       # 已加载但为空 == 零插件


# ── 1. _safe_name ───────────────────────────────────────────────

class TestSafeName(unittest.TestCase):

    def test_keeps_contact_name(self):
        # 修复前：任何输入都返回 "_"
        self.assertEqual(_safe_name("2427班级群（野生）"), "2427班级群（野生）")
        self.assertEqual(_safe_name("高途思维海超老师"), "高途思维海超老师")
        self.assertEqual(_safe_name("文件传输助手"), "文件传输助手")

    def test_replaces_illegal_chars(self):
        self.assertEqual(_safe_name('a<b>c:d"e/f\\g|h?i*j'), "a_b_c_d_e_f_g_h_i_j")

    def test_strips_and_falls_back(self):
        self.assertEqual(_safe_name("  带空格  "), "带空格")
        self.assertEqual(_safe_name("结尾有点..."), "结尾有点")
        self.assertEqual(_safe_name(""), "chat")
        self.assertEqual(_safe_name(None), "chat")
        self.assertEqual(len(_safe_name("x" * 80)), 48)

    def test_windows_reserved(self):
        self.assertEqual(_safe_name("CON"), "_CON")
        self.assertEqual(_safe_name("nul.txt"), "_nul.txt")


# ── 2. CDATA 解析 ───────────────────────────────────────────────

class TestCdata(unittest.TestCase):

    def test_cdata_content_is_preserved(self):
        """修复前 CDATA 会被替换成一个 0x01 控制字符。"""
        got = api_chat._xml_text("<![CDATA[标题内容]]>")
        self.assertEqual(got, "标题内容")
        self.assertNotIn("\x01", got or "")

    def test_appmsg_title_from_cdata(self):
        xml = ('<appmsg><title><![CDATA[一个链接标题]]></title>'
               '<url><![CDATA[https://example.com/x]]></url></appmsg>')
        title, url, _des = api_chat._parse_appmsg(xml)
        self.assertEqual(title, "一个链接标题")
        self.assertEqual(url, "https://example.com/x")


# ── 3. 分片索引 ─────────────────────────────────────────────────

class TestShardIndex(TempRootCase):

    def test_index_finds_only_real_shards(self):
        acc, _account, chat = make_account(self.tmp, n_texts=3)
        shards = api_chat.shards_for(acc, chat)
        self.assertEqual([p.name for p in shards], ["message_0.db"])
        # 不含 Msg_ 表的库不应出现在索引里
        idx = api_chat.shard_index(acc / "message")
        self.assertNotIn("media_0.db", [p.name for ps in idx.values() for p in ps])

    def test_unknown_chat_returns_empty(self):
        acc, _account, _chat = make_account(self.tmp)
        self.assertEqual(api_chat.shards_for(acc, "wxid_nobody"), [])

    def test_index_is_cached(self):
        acc, _account, chat = make_account(self.tmp)
        api_chat.shards_for(acc, chat)
        first = api_chat.shard_index(acc / "message")
        second = api_chat.shard_index(acc / "message")
        self.assertIs(first, second, "第二次调用应命中缓存（同一对象）")

    def test_message_tables_by_shard(self):
        acc, _account, chat = make_account(self.tmp)
        by_shard = api_chat.message_tables_by_shard(acc)
        only = acc / "message" / "message_0.db"
        self.assertEqual(len(by_shard), 1)
        self.assertIn(only, by_shard)
        self.assertEqual(by_shard[only], [_msg_table(chat)])


# ── 4. message_stream / count_messages ──────────────────────────

class TestSessionsApi(TempRootCase):

    def test_filters_ghost_sessions_and_marks_official_accounts(self):
        acc, account, chat = make_account(self.tmp, n_texts=3)
        sdb = acc / "session" / "session.db"
        conn = sqlite3.connect(sdb)
        conn.executemany("INSERT INTO SessionTable VALUES (?,?,?)", [
            ("gh_live", "公众号摘要", 1_700_000_100),
            ("gh_empty", "", 0),
            ("brandsessionholder", "聚合入口", 1_700_000_200),
            ("@placeholder_foldgroup", "占位入口", 1_700_000_201),
        ])
        conn.commit(); conn.close()
        cdb = acc / "contact" / "contact.db"
        conn = sqlite3.connect(cdb)
        conn.execute("INSERT INTO contact VALUES (?,?,?,?)",
                     ("gh_live", "公众号A", "", ""))
        conn.commit(); conn.close()

        from siwx.server import app
        data = app.test_client().get(f"/api/chat/sessions?account={account}").get_json()
        usernames = {s["username"]: s for s in data["sessions"]}
        self.assertIn(chat, usernames)
        self.assertIn("gh_live", usernames)
        self.assertTrue(usernames["gh_live"]["is_official"])
        self.assertEqual(usernames["gh_live"]["kind"], "official")
        self.assertEqual(usernames["gh_live"]["display"], "公众号A")
        self.assertNotIn("gh_empty", usernames)
        self.assertNotIn("brandsessionholder", usernames)
        self.assertNotIn("@placeholder_foldgroup", usernames)

    def test_sessions_api_uses_cache_after_first_call(self):
        acc, account, _chat = make_account(self.tmp, n_texts=3)
        from siwx.server import app
        client = app.test_client()
        self.assertEqual(client.get(f"/api/chat/sessions?account={account}").status_code, 200)

        opened = []
        real_connect = sqlite3.connect
        def spy(path, *a, **k):
            opened.append(Path(path).name)
            return real_connect(path, *a, **k)
        sqlite3.connect = spy
        try:
            r = client.get(f"/api/chat/sessions?account={account}")
        finally:
            sqlite3.connect = real_connect
        self.assertEqual(r.status_code, 200)
        self.assertEqual(opened, [], "会话列表缓存命中时不应再打开 contact/session 数据库")


class TestMessageStream(TempRootCase):

    def test_stream_yields_all_in_order(self):
        acc, account, chat = make_account(self.tmp, n_texts=12)
        from siwx.export_stream import message_stream
        msgs = list(message_stream(acc, chat, account=account))
        self.assertEqual(len(msgs), 12)
        ts = [m["createTime"] for m in msgs]
        self.assertEqual(ts, sorted(ts), "必须按时间正序")
        self.assertEqual(msgs[0]["content"], "第 0 条消息")
        self.assertEqual(msgs[0]["senderDisplayName"], "测试好友")

    def test_count_matches_stream(self):
        acc, account, chat = make_account(self.tmp, n_texts=7)
        from siwx.export_stream import count_messages, message_stream
        self.assertEqual(count_messages(acc, chat),
                         len(list(message_stream(acc, chat, account=account))))

    def test_shard_scan_is_avoided(self):
        """索引建好之后，不含 Msg_ 表的库不应再被打开。

        索引本身需要扫一遍目录（这是必要的一次性成本）；收益体现在后续调用：
        每个会话的两遍导出、多会话批量、聊天页都直接命中缓存。
        """
        acc, account, chat = make_account(self.tmp, n_texts=3)
        from siwx.export_stream import message_stream

        # 预热：第一次会扫全部 *.db 建立索引
        list(message_stream(acc, chat, account=account))

        opened = []
        real_connect = sqlite3.connect

        def spy(path, *a, **k):
            opened.append(Path(path).name)
            return real_connect(path, *a, **k)

        sqlite3.connect = spy
        try:
            list(message_stream(acc, chat, account=account))
        finally:
            sqlite3.connect = real_connect

        self.assertNotIn("media_0.db", opened, "索引未生效：仍在打开无关分片")
        self.assertNotIn("message_fts.db", opened, "索引未生效：仍在打开无关分片")
        self.assertIn("message_0.db", opened)

    def test_index_survives_dir_change(self):
        """目录内容变化（mtime 改变）后索引应自动失效并重建。"""
        acc, account, chat = make_account(self.tmp, n_texts=3)
        self.assertEqual(len(api_chat.shards_for(acc, chat)), 1)
        time.sleep(0.01)
        make_shard(acc / "message" / "message_1.db", chat, ["后加的"])
        self.assertEqual(len(api_chat.shards_for(acc, chat)), 2)


# ── 5. 导出：命名、HTML、JSON 干净性 ────────────────────────────

class TestExport(TempRootCase):

    def _export(self, fmt, **kw):
        acc, account, chat = self.acc, self.account, self.chat
        return exporter.run_export(
            acc, account, chat, self.display, fmt,
            want_messages=True, want_media=kw.pop("media", False),
            want_avatars=kw.pop("avatars", False),
            export_root=self.tmp / "exports", pack="none",
            progress=lambda p, m: None)

    def setUp(self):
        super().setUp()
        self.acc, self.account, self.chat = make_account(self.tmp, n_texts=12)
        self.display = "测试会话名ABC"

    def test_all_formats_embed_display_name(self):
        for fmt in ("json", "html", "txt", "csv", "markdown",
                    "toml", "sqlite", "xlsx"):
            with self.subTest(fmt=fmt):
                res = self._export(fmt)
                self.assertIn(self.display, Path(res["export_dir"]).name)
                self.assertIn(self.display, Path(res["file"]).name)
                self.assertTrue(Path(res["file"]).is_file())
                self.assertEqual(res["message_count"], 12)

    def test_html_export_not_broken(self):
        """修复前会抛 KeyError: 'displayName'。"""
        res = self._export("html")
        html = Path(res["file"]).read_text(encoding="utf-8")
        self.assertIn("window.CHAT_DATA", html)
        i = html.index("window.CHAT_DATA = ") + len("window.CHAT_DATA = ")
        data, _ = json.JSONDecoder().raw_decode(html[i:])
        meta = data["meta"]
        self.assertEqual(meta["sessionName"], self.display)
        self.assertEqual(meta["sessionId"], self.chat)
        self.assertGreater(meta["dateRange"]["start"], 0)
        self.assertGreater(meta["dateRange"]["end"], 0)
        self.assertEqual(meta["messageCount"], len(data["messages"]))

    def test_json_session_has_no_internal_keys(self):
        res = self._export("json")
        d = json.loads(Path(res["file"]).read_text(encoding="utf-8"))
        self.assertNotIn("_avatar_map", d["session"])
        for k in d["session"]:
            self.assertFalse(k.startswith("_"), f"内部键泄漏: {k}")

    def test_multi_export_creates_separate_dirs(self):
        """修复前所有会话都写进同一个 "_" 目录。"""
        acc, account, chat = self.acc, self.account, self.chat
        res = exporter.run_export_multi(
            acc, account,
            [{"chat": chat, "display": "会话甲"}, {"chat": chat, "display": "会话乙"}],
            fmt="json", want_media=False, want_avatars=False,
            export_root=self.tmp / "multi", pack="folder")
        self.assertEqual(res["ok_count"], 2)
        names = sorted(p.name for p in Path(res["total_dir"]).iterdir() if p.is_dir())
        self.assertEqual(len(names), 2, f"目录未按会话隔离: {names}")


# ── 6. 媒体：扩展名 + chat/ts 透传 ──────────────────────────────

class TestVoiceMedia(TempRootCase):

    def test_parse_voice_meta_and_read_silk_blob(self):
        from siwx import voice
        acc, account, chat = make_account(self.tmp, n_texts=1)
        db = acc / "message" / "media_0.db"
        conn = sqlite3.connect(db)
        conn.execute("DROP TABLE unrelated")
        conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
        conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
        conn.execute("CREATE TABLE VoiceInfo (chat_name_id INTEGER, create_time INTEGER, "
                     "local_id INTEGER, svr_id INTEGER, voice_data BLOB, data_index TEXT)")
        raw = b"\x02#!SILK_V3\x00\x01voice"
        conn.execute("INSERT INTO VoiceInfo VALUES (?,?,?,?,?,?)",
                     (1, 1700000000, 9, 123456789, raw, "0"))
        conn.commit(); conn.close()
        meta = voice.parse_voice_meta('<msg><voicemsg voicelength="2429" length="3990" voiceformat="4" /></msg>')
        self.assertEqual(meta["durationMs"], 2429)
        body, info = voice.get_voice(acc, chat=chat, local_id=9, svr_id=123456789, ts=1700000000)
        self.assertEqual(body, b"#!SILK_V3\x00\x01voice")
        self.assertEqual(info["silkOffset"], 1)

    def test_voice_api_serves_silk(self):
        acc, account, chat = make_account(self.tmp, n_texts=1)
        db = acc / "message" / "media_0.db"
        conn = sqlite3.connect(db)
        conn.execute("DROP TABLE unrelated")
        conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
        conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
        conn.execute("CREATE TABLE VoiceInfo (chat_name_id INTEGER, create_time INTEGER, "
                     "local_id INTEGER, svr_id INTEGER, voice_data BLOB, data_index TEXT)")
        conn.execute("INSERT INTO VoiceInfo VALUES (?,?,?,?,?,?)",
                     (1, 1700000000, 9, 123456789, b"\x02#!SILK_V3abc", "0"))
        conn.commit(); conn.close()
        from siwx.server import app
        r = app.test_client().get(
            f"/api/chat/media/voice?account={account}&chat={chat}&local_id=9&svr_id=123456789&ts=1700000000")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data, b"#!SILK_V3abc")
        self.assertEqual(r.headers.get("X-SIWX-Voice-Format"), "silk")

    def test_pcm_to_wav_uses_stdlib_container(self):
        from siwx import voice
        wav = voice.pcm_to_wav(b"\x00\x00\x01\x00", sample_rate=24000)
        self.assertTrue(wav.startswith(b"RIFF"))
        self.assertIn(b"WAVE", wav[:16])
        self.assertGreater(len(wav), 44)

    def test_transcode_prefers_pilk_backend(self):
        from siwx import voice
        seen = []
        old_pilk = voice._decode_silk_to_pcm_with_pilk
        old_cmd = voice._decode_silk_to_pcm_with_command
        try:
            voice._decode_silk_to_pcm_with_pilk = lambda data: (seen.append("pilk") or b"\x00\x00", "", "pilk")
            voice._decode_silk_to_pcm_with_command = lambda data, rate: (seen.append("cmd") or b"\x01\x00", "", "cmd")
            body, meta = voice.transcode_voice(b"#!SILK_V3abc", "wav")
        finally:
            voice._decode_silk_to_pcm_with_pilk = old_pilk
            voice._decode_silk_to_pcm_with_command = old_cmd
        self.assertEqual(seen, ["pilk"])
        self.assertTrue(body.startswith(b"RIFF"))
        self.assertEqual(meta["engine"], "pilk")

    def test_bundled_decoder_path_is_available_as_fallback(self):
        from siwx import voice
        vendor = self.tmp / "siwx" / "vendor" / "silk-decoder" / "windows"
        vendor.mkdir(parents=True, exist_ok=True)
        exe = vendor / ("silk_v3_decoder.exe" if os.name == "nt" else "silk_v3_decoder")
        exe.write_bytes(b"fake")
        old_roots = voice._resource_roots
        try:
            voice._resource_roots = lambda: [self.tmp / "siwx"]
            candidates = voice._decoder_candidates()
        finally:
            voice._resource_roots = old_roots
        self.assertTrue(candidates)
        self.assertEqual(candidates[0][1], [str(exe)])
        self.assertTrue(candidates[0][0].startswith("bundled:"))

    def test_voice_api_transcodes_wav_when_decoder_available(self):
        acc, account, chat = make_account(self.tmp, n_texts=1)
        db = acc / "message" / "media_0.db"
        conn = sqlite3.connect(db)
        conn.execute("DROP TABLE unrelated")
        conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
        conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
        conn.execute("CREATE TABLE VoiceInfo (chat_name_id INTEGER, create_time INTEGER, "
                     "local_id INTEGER, svr_id INTEGER, voice_data BLOB, data_index TEXT)")
        conn.execute("INSERT INTO VoiceInfo VALUES (?,?,?,?,?,?)",
                     (1, 1700000000, 9, 123456789, b"\x02#!SILK_V3abc", "0"))
        conn.commit(); conn.close()
        from siwx import voice
        from siwx.server import app
        old = voice.transcode_voice
        try:
            voice.transcode_voice = lambda data, target="wav": (
                b"RIFFxxxxWAVEfmt ", {"format": "wav", "mimetype": "audio/wav", "ext": "wav", "engine": "fake"})
            r = app.test_client().get(
                f"/api/chat/media/voice?account={account}&chat={chat}&local_id=9&svr_id=123456789&ts=1700000000&format=wav")
        finally:
            voice.transcode_voice = old
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, "audio/wav")
        self.assertEqual(r.headers.get("X-SIWX-Voice-Format"), "wav")
        self.assertEqual(r.headers.get("X-SIWX-Voice-Transcoder"), "fake")


class TestMediaExport(TempRootCase):

    def setUp(self):
        super().setUp()
        self.acc, self.account, self.chat = make_account(self.tmp, n_texts=6)

    def test_extension_follows_actual_content(self):
        """修复前 media 引用恒为 .jpg，即使实际写的是 .png。"""
        from siwx import exporter as ex

        def fake_get_image(account, md5, acc_dir, **kw):
            return b"\x89PNG\r\n\x1a\n" + b"\x00" * 32, "image/png"

        old = ex.media.get_image
        ex.media.get_image = fake_get_image
        try:
            dest = self.tmp / "media_out"
            dest.mkdir(parents=True, exist_ok=True)
            out = ex._try_decrypt(str(self.acc), self.account, self.chat,
                                  "a" * 32, None, 1, 1_700_000_000,
                                  dest / "0000_aaaaaaaaaaaa.jpg")
            self.assertIsNotNone(out)
            self.assertEqual(out.suffix, ".png")
            self.assertTrue(out.is_file())
            self.assertEqual(out.name, "0000_aaaaaaaaaaaa.png")
        finally:
            ex.media.get_image = old

    def test_chat_and_ts_are_forwarded(self):
        """修复前未传 chat/ts，attach/Bubble/Thumb 三级来源全部失效。"""
        from siwx import exporter as ex
        seen = {}

        def fake_get_image(account, md5, acc_dir, **kw):
            seen.update(kw)
            return None, "nope"

        old = ex.media.get_image
        ex.media.get_image = fake_get_image
        try:
            ex._try_decrypt(str(self.acc), self.account, self.chat,
                            "b" * 32, "c" * 32, 42, 1_700_000_123,
                            self.tmp / "x.jpg")
        finally:
            ex.media.get_image = old
        self.assertEqual(seen.get("chat"), self.chat)
        self.assertEqual(seen.get("ts"), 1_700_000_123)
        self.assertEqual(seen.get("local_id"), 42)
        self.assertEqual(seen.get("bubble_md5"), "c" * 32)

    def _add_voice_message(self, local_id=99, svr_id=123456789, ts=1_700_000_099):
        t = _msg_table(self.chat)
        conn = sqlite3.connect(self.acc / "message" / "message_0.db")
        conn.execute(f"INSERT INTO [{t}] VALUES (?,?,?,?,?,?,?,?)",
                     (local_id, svr_id, 34, ts, 0, 1,
                      b'<msg><voicemsg voicelength="1000" length="12" /></msg>', None))
        conn.commit(); conn.close()

        db = self.acc / "message" / "media_0.db"
        conn = sqlite3.connect(db)
        conn.execute("DROP TABLE unrelated")
        conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
        conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (self.chat,))
        conn.execute("CREATE TABLE VoiceInfo (chat_name_id INTEGER, create_time INTEGER, "
                     "local_id INTEGER, svr_id INTEGER, voice_data BLOB, data_index TEXT)")
        conn.execute("INSERT INTO VoiceInfo VALUES (?,?,?,?,?,?)",
                     (1, ts, local_id, svr_id, b"\x02#!SILK_V3abc", "0"))
        conn.commit(); conn.close()

    def test_export_includes_transcoded_voice_media(self):
        from siwx import exporter as ex
        self._add_voice_message()
        old = ex.voice.transcode_voice
        try:
            ex.voice.transcode_voice = lambda data, target="wav": (
                b"RIFFxxxxWAVEfmt ", {"format": "wav", "mimetype": "audio/wav", "ext": "wav", "engine": "fake"})
            res = ex.run_export(self.acc, self.account, self.chat, "测试好友",
                                fmt="html", want_media=False, want_voice=True,
                                want_avatars=False,
                                export_root=self.tmp / "exports", pack="none")
        finally:
            ex.voice.transcode_voice = old
        html = Path(res["file"]).read_text(encoding="utf-8")
        self.assertEqual(res["voice_count"], 1)
        self.assertTrue((Path(res["file"]).parent / "media" / "voice_0000_99.wav").is_file())
        self.assertIn("voice_0000_99.wav", html)

    def test_all_formats_can_reference_exported_voice(self):
        from siwx import exporter as ex
        self._add_voice_message(local_id=77, svr_id=777, ts=1_700_000_077)
        old = ex.voice.transcode_voice
        try:
            ex.voice.transcode_voice = lambda data, target="wav": (
                b"RIFFxxxxWAVEfmt ", {"format": "wav", "mimetype": "audio/wav", "ext": "wav", "engine": "fake"})
            for fmt in ("json", "html", "txt", "csv", "markdown", "toml", "sqlite", "xlsx"):
                with self.subTest(fmt=fmt):
                    res = ex.run_export(self.acc, self.account, self.chat, "测试好友",
                                        fmt=fmt, want_media=False, want_voice=True,
                                        want_avatars=False, export_root=self.tmp / "voice_formats",
                                        folder_name=f"voice_{fmt}", pack="none")
                    out_file = Path(res["file"])
                    voice_file = out_file.parent / "media" / "voice_0000_77.wav"
                    self.assertTrue(voice_file.is_file())
                    self.assertEqual(res["voice_count"], 1)
                    if fmt == "sqlite":
                        conn = sqlite3.connect(out_file)
                        vals = [r[0] for r in conn.execute("SELECT mediaFile FROM messages WHERE mediaFile IS NOT NULL")]
                        conn.close()
                        self.assertIn("media/voice_0000_77.wav", vals)
                    elif fmt == "xlsx":
                        from openpyxl import load_workbook
                        wb = load_workbook(out_file, read_only=True)
                        vals = [row[-1] for row in wb.active.iter_rows(values_only=True)]
                        self.assertIn("media/voice_0000_77.wav", vals)
                    else:
                        self.assertIn("voice_0000_77.wav", out_file.read_text(encoding="utf-8"))
        finally:
            ex.voice.transcode_voice = old


# ── 7. messages 的 limit 下限 ───────────────────────────────────

class TestAvatarApi(TempRootCase):

    def test_owner_avatar_falls_back_to_clean_wxid(self):
        """输出目录名可能带 _数字后缀，但头像库里本人是原始 wxid。"""
        acc, account, _chat = make_account(self.tmp, account="wxid_owner_1234", n_texts=1)
        (acc / "head_image").mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(acc / "head_image" / "head_image.db")
        conn.execute("CREATE TABLE head_image (username TEXT PRIMARY KEY, md5 TEXT, image_buffer BLOB, update_time INTEGER)")
        conn.execute("INSERT INTO head_image VALUES (?,?,?,?)",
                     ("wxid_owner", "m", b"JPEGDATA", 1))
        conn.commit(); conn.close()
        from siwx.server import app
        r = app.test_client().get("/api/chat/avatar?account=wxid_owner_1234&username=wxid_owner_1234")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data, b"JPEGDATA")


class TestSettingsAutoSync(TempRootCase):

    def test_auto_sync_settings_roundtrip(self):
        from siwx.server import app
        c = app.test_client()
        r = c.post("/api/settings/auto-sync", json={"enabled": True, "interval_minutes": 5})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["enabled"])
        self.assertEqual(r.get_json()["interval_minutes"], 5)
        r2 = c.get("/api/settings/auto-sync")
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(r2.get_json()["interval_minutes"], 5)

    def test_auto_sync_interval_is_clamped(self):
        from siwx.server import app
        r = app.test_client().post("/api/settings/auto-sync", json={"enabled": True, "interval_minutes": 99999})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["interval_minutes"], 1440)


class TestStatsApi(TempRootCase):
    """聊天统计：跨分片聚合、类型分布、时间维度、缓存与过滤。"""

    @staticmethod
    def _make_stats_account(root: Path, account="wxid_stats"):
        """构造含多分片、多类型、多发送者的统计样本。"""
        acc = root / "output" / account
        msg_dir = acc / "message"
        msg_dir.mkdir(parents=True, exist_ok=True)
        # 固定基准时间，便于断言小时/月份（本地时区，用正午避开跨日边界）
        base = 1_700_000_000

        def shard(name, chat, rows):
            conn = sqlite3.connect(msg_dir / name)
            t = _msg_table(chat)
            conn.execute(f"""CREATE TABLE [{t}] (
                local_id INTEGER PRIMARY KEY, server_id INTEGER, local_type INTEGER,
                create_time INTEGER, origin_source INTEGER, real_sender_id INTEGER,
                message_content BLOB, packed_info_data BLOB)""")
            for i, (ltype, ts, sid) in enumerate(rows):
                conn.execute(f"INSERT INTO [{t}] VALUES (?,?,?,?,?,?,?,?)",
                             (i + 1, 100 + i, ltype, ts, 0, sid, b"x", None))
            conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
            conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
            conn.commit()
            conn.close()

        # 分片 0：文本 + 图片，发送者 1
        shard("message_0.db", "wxid_a",
              [(1, base, 1), (1, base + 3600, 1), (3, base + 7200, 1)])
        # 分片 1：表情 + 语音，发送者 2
        shard("message_1.db", "wxid_b",
              [(47, base + 100, 2), (34, base + 200, 2)])
        # 分片 2：系统消息 + 一年前的文本
        shard("message_2.db", "wxid_c",
              [(10000, base + 300, 0), (1, base - 400 * 86400, 1)])

        # 三个不同会话 → chat_count 应为 3；联系人库用于验证排行显示昵称。
        (acc / "contact").mkdir(parents=True, exist_ok=True)
        c = sqlite3.connect(acc / "contact" / "contact.db")
        c.execute("CREATE TABLE contact (username TEXT, remark TEXT, nick_name TEXT, "
                  "alias TEXT, verify_flag INTEGER)")
        c.executemany("INSERT INTO contact VALUES (?,?,?,?,?)", [
            ("wxid_a", "好友A", "", "", 0),
            ("wxid_b", "", "好友B", "", 0),
            ("wxid_c", "", "好友C", "", 0),
            ("gh_news", "", "公众号", "news_alias", 1053),
            ("group@chatroom", "", "群聊", "", 0),
        ])
        c.commit()
        c.close()
        return acc, account

    def test_overview_totals_and_types(self):
        self._make_stats_account(self.tmp)
        from siwx.server import app
        r = app.test_client().get("/api/stats/overview?account=wxid_stats")
        self.assertEqual(r.status_code, 200)
        d = r.get_json()
        self.assertEqual(d["total"], 7)
        self.assertEqual(d["chat_count"], 3)
        self.assertEqual(d["shards"], 3)
        # 类型分组：文本 3、图片 1、表情 1、语音 1、系统 1（合计 7）
        groups = {g["label"]: g["count"] for g in d["type_groups"]}
        self.assertEqual(groups.get("文本"), 3)
        self.assertEqual(groups.get("图片"), 1)
        self.assertEqual(groups.get("表情"), 1)
        self.assertEqual(groups.get("语音"), 1)
        self.assertEqual(groups.get("系统"), 1)
        # 分组之和必须等于总量，否则图表会缺数据
        self.assertEqual(sum(groups.values()), d["total"])

    def test_hour_and_weekday_histograms(self):
        self._make_stats_account(self.tmp)
        from siwx.server import app
        d = app.test_client().get("/api/stats/overview?account=wxid_stats").get_json()
        self.assertEqual(len(d["by_hour"]), 24)
        self.assertEqual(len(d["by_weekday"]), 7)
        # 直方图总量应等于消息总数（每条消息恰好落进一个小格）
        self.assertEqual(sum(d["by_hour"]), d["total"])
        self.assertEqual(sum(d["by_weekday"]), d["total"])

    def test_month_series_spans_multiple_months(self):
        self._make_stats_account(self.tmp)
        from siwx.server import app
        d = app.test_client().get("/api/stats/overview?account=wxid_stats").get_json()
        # 样本含一年前的消息 → 至少两个不同月份
        self.assertGreaterEqual(len(d["by_month"]), 2)
        self.assertEqual(sum(m["count"] for m in d["by_month"]), d["total"])

    def test_top_senders_only_private_and_resolves_nickname(self):
        acc, _account = self._make_stats_account(self.tmp)
        # 群聊与公众号给更多消息，若未过滤会排到第一。
        msg_dir = acc / "message"
        def add_shard(name, chat, n):
            conn = sqlite3.connect(msg_dir / name)
            t = _msg_table(chat)
            conn.execute(f"""CREATE TABLE [{t}] (
                local_id INTEGER PRIMARY KEY, server_id INTEGER, local_type INTEGER,
                create_time INTEGER, origin_source INTEGER, real_sender_id INTEGER,
                message_content BLOB, packed_info_data BLOB)""")
            for i in range(n):
                conn.execute(f"INSERT INTO [{t}] VALUES (?,?,?,?,?,?,?,?)",
                             (i + 1, 200 + i, 1, 1_700_000_000 + i, 0, 1, b"x", None))
            conn.execute("CREATE TABLE Name2Id (user_name TEXT)")
            conn.execute("INSERT INTO Name2Id(rowid, user_name) VALUES (1, ?)", (chat,))
            conn.commit()
            conn.close()
        add_shard("message_3.db", "group@chatroom", 20)
        add_shard("message_4.db", "gh_news", 30)

        from siwx import stats
        stats.clear_cache()
        from siwx.server import app
        d = app.test_client().get("/api/stats/overview?account=wxid_stats&refresh=1").get_json()
        top = {s["wxid"]: s for s in d["top_senders"]}
        self.assertNotIn("group@chatroom", top)
        self.assertNotIn("gh_news", top)
        # 排行按私聊会话聚合，且展示联系人备注/昵称。
        self.assertEqual(top["wxid_a"]["count"], 3)
        self.assertEqual(top["wxid_a"]["name"], "好友A")
        self.assertEqual(top["wxid_b"]["name"], "好友B")

    def test_date_filter_narrows_all_statistics(self):
        self._make_stats_account(self.tmp)
        from siwx.server import app
        c = app.test_client()
        full = c.get("/api/stats/overview?account=wxid_stats").get_json()
        filtered = c.get(
            "/api/stats/overview?account=wxid_stats&start=2023-11-15&end=2023-11-15"
        ).get_json()
        # 样本里 6 条在 2023-11-15，1 条在 400 天前；过滤应影响整页指标。
        self.assertEqual(full["total"], 7)
        self.assertEqual(filtered["total"], 6)
        self.assertEqual(sum(m["count"] for m in filtered["by_month"]), 6)
        self.assertEqual(sum(filtered["by_hour"]), 6)
        self.assertEqual(sum(filtered["by_weekday"]), 6)
        groups = {g["label"]: g["count"] for g in filtered["type_groups"]}
        self.assertEqual(groups.get("文本"), 2)
        top = {s["wxid"]: s["count"] for s in filtered["top_senders"]}
        # wxid_c 在当天只有 1 条系统消息；一年前那条文本不应混进范围内。
        self.assertEqual(top.get("wxid_c"), 1)

    def test_cache_hit_after_first_compute(self):
        acc, account = self._make_stats_account(self.tmp)
        from siwx import stats
        stats.clear_cache()
        stats.compute_stats(account)
        self.assertTrue((acc / ".siwx_stats.json").is_file())
        sig = stats.signature(account)
        self.assertIsNotNone(sig)
        # 磁盘缓存可被读取（内容与签名匹配）
        cached = stats._load_disk_cache(account, sig)
        self.assertIsNotNone(cached)
        self.assertEqual(cached["total"], 7)

    def test_cache_invalidated_when_shard_changes(self):
        acc, account = self._make_stats_account(self.tmp)
        from siwx import stats
        stats.clear_cache()
        stats.compute_stats(account)
        old_sig = stats.signature(account)
        # 追加一个分片 → 签名必须变化，否则统计会永久停在旧结果
        conn = sqlite3.connect(acc / "message" / "message_9.db")
        t = _msg_table("wxid_new")
        conn.execute(f"""CREATE TABLE [{t}] (
            local_id INTEGER PRIMARY KEY, server_id INTEGER, local_type INTEGER,
            create_time INTEGER, origin_source INTEGER, real_sender_id INTEGER,
            message_content BLOB, packed_info_data BLOB)""")
        conn.execute(f"INSERT INTO [{t}] VALUES (1,1,1,1700000500,0,1,?,NULL)", (b"x",))
        conn.commit()
        conn.close()
        new_sig = stats.signature(account)
        self.assertNotEqual(old_sig, new_sig)
        self.assertIsNone(stats._load_disk_cache(account, new_sig))
        self.assertEqual(stats.compute_stats(account)["total"], 8)

    def test_accounts_endpoint_lists_only_decrypted(self):
        self._make_stats_account(self.tmp)
        (self.tmp / "output" / "wxid_no_msg").mkdir(parents=True, exist_ok=True)
        from siwx.server import app
        d = app.test_client().get("/api/stats/accounts").get_json()
        names = [a["wxid"] for a in d["accounts"]]
        self.assertIn("wxid_stats", names)
        self.assertNotIn("wxid_no_msg", names)

    def test_overview_requires_account(self):
        from siwx.server import app
        self.assertEqual(app.test_client().get("/api/stats/overview").status_code, 400)

    def test_overview_unknown_account_is_404(self):
        from siwx.server import app
        r = app.test_client().get("/api/stats/overview?account=wxid_missing")
        self.assertEqual(r.status_code, 404)

    def test_refresh_endpoint_recomputes(self):
        self._make_stats_account(self.tmp)
        from siwx.server import app
        r = app.test_client().post("/api/stats/refresh", json={"account": "wxid_stats"})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])
        self.assertEqual(r.get_json()["total"], 7)

    def test_empty_account_returns_zeroes(self):
        acc = self.tmp / "output" / "wxid_empty"
        (acc / "message").mkdir(parents=True, exist_ok=True)
        from siwx.server import app
        d = app.test_client().get("/api/stats/overview?account=wxid_empty").get_json()
        self.assertEqual(d["total"], 0)
        self.assertEqual(d["chat_count"], 0)
        self.assertEqual(d["type_groups"], [])


class TestLimitGuard(TempRootCase):

    def test_negative_limit_is_clamped(self):
        acc, account, chat = make_account(self.tmp, n_texts=30)
        from siwx.server import app
        c = app.test_client()
        r = c.get(f"/api/chat/messages?account={account}&chat={chat}&limit=-1")
        self.assertEqual(r.status_code, 200)
        n = len(r.get_json()["messages"])
        self.assertLessEqual(n, 300)
        self.assertGreater(n, 0)


class TestChatTimelineAndStats(TempRootCase):

    def test_timeline_and_conversation_stats(self):
        _acc, account, chat = make_account(self.tmp, n_texts=6)
        from siwx.server import app
        c = app.test_client()
        tl = c.get(f"/api/chat/timeline?account={account}&chat={chat}")
        self.assertEqual(tl.status_code, 200)
        self.assertEqual(tl.get_json()["total"], 6)
        self.assertTrue(tl.get_json()["months"])
        month = tl.get_json()["months"][0]["month"]
        days = c.get(f"/api/chat/timeline?account={account}&chat={chat}&month={month}")
        self.assertEqual(days.status_code, 200)
        self.assertEqual(days.get_json()["total"], 6)
        self.assertTrue(days.get_json()["days"])

        st = c.get(f"/api/chat/stats?account={account}&chat={chat}")
        self.assertEqual(st.status_code, 200)
        d = st.get_json()
        self.assertEqual(d["total"], 6)
        self.assertEqual(d["types"][0]["label"], "文本消息")
        self.assertGreaterEqual(d["active_days"], 1)


class TestManualWechatPaths(TempRootCase):

    def test_validate_endpoint_persists_manual_path(self):
        db_dir = self.tmp / "custom" / "wxid_manual" / "db_storage"
        db_dir.mkdir(parents=True)
        from siwx.server import app
        from siwx.discover import load_manual_data_dirs, manual_paths_file
        r = app.test_client().post("/api/discover/validate", json={"path": str(db_dir)})
        self.assertEqual(r.status_code, 200)
        d = r.get_json()
        self.assertTrue(d["ok"])
        self.assertTrue(d["saved"])
        self.assertTrue(manual_paths_file().is_file())
        self.assertIn(("wxid_manual", str(db_dir.resolve())), load_manual_data_dirs())

    def test_resolves_account_subdir_and_database_file(self):
        from siwx.discover import resolve_db_paths
        db_dir = self.tmp / "xwechat_files" / "wxid_one" / "db_storage"
        message = db_dir / "message"
        message.mkdir(parents=True)
        db_file = message / "message_0.db"
        db_file.write_bytes(b"")
        expected = [{"wxid": "wxid_one", "db_dir": str(db_dir.resolve())}]
        self.assertEqual(resolve_db_paths(str(db_dir.parent)), expected)
        self.assertEqual(resolve_db_paths(str(message)), expected)
        self.assertEqual(resolve_db_paths(str(db_file)), expected)

    def test_xwechat_root_finds_multiple_accounts(self):
        from siwx.discover import validate_db_path
        root = self.tmp / "xwechat_files"
        for wxid in ("wxid_a", "wxid_b"):
            (root / wxid / "db_storage").mkdir(parents=True)
        d = validate_db_path(str(root))
        self.assertTrue(d["ok"])
        self.assertEqual(d["account_count"], 2)
        self.assertEqual({a["wxid"] for a in d["accounts"]}, {"wxid_a", "wxid_b"})


# ── 8. 解密原子写 ───────────────────────────────────────────────

class TestDecryptAtomic(unittest.TestCase):
    """构造一个合法的 SQLCipher 4 单页库，验证解密与失败时的原子性。"""

    @staticmethod
    def _encrypt_page(plain_body: bytes, pageno: int, enc_key: bytes,
                      salt: bytes, is_first: bool):
        from Crypto.Cipher import AES
        mac_salt = bytes(b ^ 0x3A for b in salt)
        mac_key = hashlib.pbkdf2_hmac("sha512", enc_key, mac_salt, 2, dklen=32)
        iv = bytes((pageno * 7 + i) & 0xFF for i in range(16))
        ct = AES.new(enc_key, AES.MODE_CBC, iv).encrypt(plain_body)
        if is_first:
            page = salt + ct + iv
        else:
            page = ct + iv
        # 与 verify_enc_key 对齐：HMAC 覆盖 page1[16:]（页 1 跳过 salt）
        mac_input = page[len(salt):] if is_first else page
        mac = hmac_mod.new(mac_key, mac_input, hashlib.sha512)
        mac.update(struct.pack("<I", pageno))
        return page + mac.digest()

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="siwx_db_"))
        from Crypto.Cipher import AES
        from siwx.sqlcipher import PAGE_SZ, RESERVE_SZ, SALT_SZ
        self.enc_key = bytes(range(32))
        salt = bytes(range(16, 32))
        body_len = PAGE_SZ - RESERVE_SZ          # 4016
        # 页 1：正文 4000 字节（salt 占掉 16）
        page1 = self._encrypt_page(bytes((i * 3) & 0xFF for i in range(body_len - SALT_SZ)),
                                   1, self.enc_key, salt, True)
        # 页 2
        page2 = self._encrypt_page(bytes((i * 5) & 0xFF for i in range(body_len)),
                                   2, self.enc_key, salt, False)
        self.src = self.tmp / "src.db"
        self.src.write_bytes(page1 + page2)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_decrypts_and_leaves_no_residue(self):
        from siwx.sqlcipher import decrypt_database
        dst = self.tmp / "out" / "dst.db"
        pages = decrypt_database(self.src, dst, self.enc_key)
        self.assertEqual(pages, 2)
        self.assertTrue(dst.is_file())
        self.assertEqual(dst.stat().st_size, 2 * 4096)
        with dst.open("rb") as f:
            self.assertEqual(f.read(16), b"SQLite format 3\x00")
        residue = list(dst.parent.glob("*.part")) + list(dst.parent.glob("*.tmp"))
        self.assertEqual(residue, [], f"残留临时文件: {residue}")

    def test_failure_does_not_clobber_existing_file(self):
        """源库密钥错误时，已存在的明文库必须保持原样。"""
        from siwx.sqlcipher import decrypt_database
        dst = self.tmp / "out" / "dst.db"
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_bytes(b"ORIGINAL-GOOD-CONTENT")
        with self.assertRaises(ValueError):
            decrypt_database(self.src, dst, bytes(32))   # 错误密钥
        self.assertEqual(dst.read_bytes(), b"ORIGINAL-GOOD-CONTENT")
        residue = list(dst.parent.glob("*.part"))
        self.assertEqual(residue, [], f"残留临时文件: {residue}")


# ── 9. CLI --json ───────────────────────────────────────────────

class TestLogsApi(unittest.TestCase):
    """/api/logs 条目是**混合形状**的：

    旧/文件日志 → [ts_ms, text]；结构化日志（siwx.logger）→ [ts_ms, level, module, text]。
    因此断言必须走统一的取文本逻辑，不能假定固定长度（插件加载等会在启动期
    就写入结构化日志，使结构化条目非空）。
    """

    @staticmethod
    def _texts(data):
        items = data.get("logs", []) if isinstance(data, dict) else (data or [])
        return [str(it[3] if len(it) >= 4 else it[1]) for it in items]

    def test_api_logs_includes_file_logger_messages(self):
        """日志页不能只看内存 ring；普通 logger 写入的文件日志也要显示。"""
        from siwx import server
        marker = f"unit-log-marker-{int(time.time() * 1000)}"
        server._siwx_logger.info(marker)
        for h in server._siwx_logger.handlers:
            try:
                h.flush()
            except Exception:
                pass
        data = server.app.test_client().get("/api/logs").get_json()
        self.assertTrue(any(marker in m for m in self._texts(data)),
                        "文件日志没有出现在 /api/logs")

    def test_404_is_not_logged_as_uncaught_error(self):
        from siwx import server
        before = len(server.app.test_client().get("/api/logs").get_json().get("logs", []))
        r = server.app.test_client().get("/__definitely_missing__")
        self.assertEqual(r.status_code, 404)
        data = server.app.test_client().get("/api/logs").get_json()
        lines = self._texts(data)
        self.assertFalse(any("__definitely_missing__" in m or "404 Not Found" in m
                             for m in lines[-20:]))
        self.assertGreaterEqual(len(lines), before)

    def test_task_exception_is_persisted_to_file_logs(self):
        from siwx import server
        marker = "unit-task-failure-marker"
        try:
            raise RuntimeError(marker)
        except Exception as e:
            server._siwx_logger.exception("任务执行失败: %s", e)
            server._flush_logs()
        data = server.app.test_client().get("/api/logs").get_json()
        self.assertTrue(any(marker in m for m in self._texts(data)),
                        "任务异常没有落盘到 /api/logs")

    def test_all_log_items_are_renderable(self):
        """混合形状下的健壮性：每条日志都能取出文本，且前端可安全渲染。"""
        from siwx import server, logger as _log
        _log.info("plugin", "unit-mixed-shape-marker")
        items = server.app.test_client().get("/api/logs").get_json()["logs"]
        self.assertTrue(items, "日志不应为空")
        texts = self._texts(items)
        self.assertTrue(all(isinstance(t, str) and t for t in texts),
                        "存在无法取文本的日志条目")
        self.assertTrue(any("unit-mixed-shape-marker" in t for t in texts))


class TestCliJson(unittest.TestCase):

    def test_json_flag_emits_parseable_json(self):
        import argparse
        from siwx import cli, extract
        old = extract.extract_all
        extract.extract_all = lambda **kw: [
            {"wxid": "wxid_x", "db_count": 1, "total_salts": 2, "verified": 2,
             "cached": 0, "duration_ms": 1, "salts": []}]
        try:
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = cli.cmd_keys_extract(
                    argparse.Namespace(json=True, no_cache=False))
            self.assertEqual(code, 0)
            self.assertEqual(json.loads(buf.getvalue())[0]["wxid"], "wxid_x")
        finally:
            extract.extract_all = old

    def test_json_flag_returns_1_when_no_accounts(self):
        import argparse
        from siwx import cli, extract
        old = extract.extract_all
        extract.extract_all = lambda **kw: []
        try:
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = cli.cmd_keys_extract(
                    argparse.Namespace(json=True, no_cache=False))
            self.assertEqual(code, 1)
        finally:
            extract.extract_all = old


# ── 10. 密码学原语未被破坏 ──────────────────────────────────────

class TestVersionSource(unittest.TestCase):

    def test_current_version_comes_from_package_init(self):
        from siwx import __version__
        from siwx.auto_update import current_version
        self.assertEqual(current_version(), __version__)
        self.assertEqual(__version__, "5.0.5")

    def test_remote_version_uses_newest_source_and_bypasses_cache(self):
        from siwx import auto_update

        seen = []

        def fake_urlopen(req, timeout):
            seen.append(req)
            version = "5.0.0" if "gh.1s.fan" in req.full_url else "5.0.1"
            return io.BytesIO(json.dumps({"version": version}).encode("utf-8"))

        with mock.patch("urllib.request.urlopen", side_effect=fake_urlopen):
            remote = auto_update.fetch_remote_version()

        self.assertEqual(remote["version"], "5.0.1")
        self.assertEqual(len(seen), len(auto_update.VERSION_URLS))
        for req in seen:
            self.assertIn("_siwx_update=", req.full_url)
            self.assertEqual(req.get_header("Cache-control"), "no-cache")
            self.assertEqual(req.get_header("Pragma"), "no-cache")

    def test_update_check_response_is_not_cached(self):
        from flask import Flask
        from siwx import api_update

        app = Flask(__name__)
        app.register_blueprint(api_update.bp)
        remote = {"version": "5.0.1", "notes": "update"}
        with mock.patch.object(api_update, "has_update",
                               return_value=(True, remote, "5.0.0")):
            response = app.test_client().get("/api/update/check")

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["has_update"])
        self.assertIn("no-store", response.headers["Cache-Control"])
        self.assertEqual(response.headers["Pragma"], "no-cache")


class TestEnvInfo(unittest.TestCase):
    """环境信息采集（供 bug 报告粘贴）：字段齐全、用户名打码、不污染 stdout。"""

    def test_collect_has_fields_required_by_issue_template(self):
        from siwx import __version__
        from siwx import env_info

        info = env_info.collect(quiet=True)
        # 这几个字段对应 issue 模板里要求用户填写的内容
        for key in ("siwx 版本", "运行模式", "操作系统", "系统版本",
                    "系统架构", "Python", "数据目录", "密钥库"):
            self.assertIn(key, info, f"缺少字段: {key}")
        self.assertEqual(info["siwx 版本"], __version__)
        self.assertIn(info["运行模式"], ("打包产物", "源码运行"))
        # 所有值都必须是可直接粘贴的字符串
        for key, value in info.items():
            self.assertIsInstance(value, str, f"{key} 不是字符串")

    def test_mask_path_hides_username_on_all_platforms(self):
        from siwx.env_info import mask_path

        cases = [
            (r"C:\Users\alice\AppData\Local\stories-in-wx", "alice"),
            ("/Users/bob/Library/Application Support/stories-in-wx", "bob"),
            ("/home/carol/.local/share/stories-in-wx", "carol"),
        ]
        for raw, user in cases:
            masked = mask_path(raw)
            self.assertNotIn(user, masked, f"用户名未打码: {masked}")
            self.assertIn("<user>", masked)

    def test_format_text_masks_current_user(self):
        from siwx import env_info

        text = env_info.format_text(quiet=True)
        self.assertIn("### 环境信息", text)
        self.assertIn(f"- siwx 版本: {env_info.__version__}", text)
        user = Path.home().name
        if user and user not in ("root",):
            self.assertNotIn(f"\\Users\\{user}", text, "Windows 用户名未打码")
            self.assertNotIn(f"/Users/{user}", text, "macOS 用户名未打码")
            self.assertNotIn(f"/home/{user}", text, "Linux 用户名未打码")

    def test_collect_quiet_writes_nothing_to_stdout(self):
        from siwx import env_info

        buf = io.StringIO()
        with redirect_stdout(buf):
            info = env_info.collect(quiet=True)
        self.assertEqual(buf.getvalue(), "", "quiet=True 仍向 stdout 输出了内容")
        self.assertTrue(info)

    def test_settings_env_api(self):
        from flask import Flask
        from siwx import __version__
        from siwx import api_settings

        app = Flask(__name__)
        app.register_blueprint(api_settings.bp)
        response = app.test_client().get("/api/settings/env")

        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertIn("info", data)
        self.assertIn("text", data)
        self.assertEqual(data["info"]["siwx 版本"], __version__)
        self.assertIn("### 环境信息", data["text"])
        # 接口返回的文本同样必须打码
        user = Path.home().name
        if user and user not in ("root",):
            self.assertNotIn(f"\\Users\\{user}", data["text"])
            self.assertNotIn(f"/Users/{user}", data["text"])

    def test_cli_doctor_prints_paste_block(self):
        import argparse
        from siwx import cli

        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cli.cmd_doctor(argparse.Namespace())

        self.assertEqual(code, 0)
        out = buf.getvalue()
        self.assertIn("### 环境信息", out)
        self.assertIn("- siwx 版本:", out)
        self.assertIn("可直接粘贴到 GitHub issue", out)


class TestContributionTemplates(unittest.TestCase):
    """贡献规范化文件存在且 YAML 合法（GitHub 表单格式错误会直接不显示）。"""

    def _parse(self, path: Path):
        try:
            import yaml
        except ImportError:
            self.skipTest("未安装 pyyaml，跳过表单格式校验")
        return yaml.safe_load(path.read_text(encoding="utf-8"))

    def test_issue_templates_present_and_valid(self):
        tpl_dir = ROOT / ".github" / "ISSUE_TEMPLATE"
        self.assertTrue(tpl_dir.is_dir(), "缺少 .github/ISSUE_TEMPLATE 目录")

        expected = {"bug_report.yml", "feature_request.yml", "question.yml"}
        found = {p.name for p in tpl_dir.glob("*.yml")}
        self.assertTrue(expected.issubset(found), f"缺少模板: {expected - found}")

        for name in sorted(expected):
            data = self._parse(tpl_dir / name)
            for key in ("name", "description", "body"):
                self.assertIn(key, data, f"{name} 缺少 {key}")
            ids = [b.get("id") for b in data["body"] if b.get("id")]
            self.assertEqual(len(ids), len(set(ids)), f"{name} 存在重复 id")
            for block in data["body"]:
                self.assertIn(block.get("type"),
                              ("markdown", "input", "textarea", "dropdown",
                               "checkboxes"),
                              f"{name} 含不支持的字段类型")

    def test_bug_report_requires_version_os_arch(self):
        """用户明确要求：bug 反馈必须填 siwx 版本、系统版本、架构。"""
        data = self._parse(ROOT / ".github" / "ISSUE_TEMPLATE" / "bug_report.yml")
        required = {b["id"] for b in data["body"]
                    if (b.get("validations") or {}).get("required")}
        for field in ("siwx_version", "os_version", "arch", "os",
                      "wechat_version", "run_mode", "env_info"):
            self.assertIn(field, required, f"bug 模板未强制要求 {field}")

    def test_bug_report_arch_is_dropdown(self):
        data = self._parse(ROOT / ".github" / "ISSUE_TEMPLATE" / "bug_report.yml")
        arch = next(b for b in data["body"] if b.get("id") == "arch")
        self.assertEqual(arch["type"], "dropdown")
        options = " ".join(arch["attributes"]["options"]).lower()
        self.assertIn("x64", options)
        self.assertIn("arm64", options)

    def test_blank_issues_disabled(self):
        data = self._parse(ROOT / ".github" / "ISSUE_TEMPLATE" / "config.yml")
        self.assertTrue(data.get("blank_issues_enabled") is False,
                        "空白 issue 应被禁用，强制走模板")

    def test_pr_template_and_contributing_present(self):
        self.assertTrue((ROOT / ".github" / "PULL_REQUEST_TEMPLATE.md").is_file())
        self.assertTrue((ROOT / "CONTRIBUTING.md").is_file())


class TestCryptoIntact(unittest.TestCase):

    def test_verify_enc_key_byte_layout(self):
        from siwx import sqlcipher as sc
        self.assertEqual(sc.PAGE_SZ - sc.RESERVE_SZ + sc.IV_SZ - sc.SALT_SZ, 4016)

    def test_handwritten_cbc_matches_stdlib(self):
        from Crypto.Cipher import AES
        from siwx.sqlcipher import PAGE_SZ, RESERVE_SZ, IV_SZ
        key, iv = bytes(range(32)), bytes(range(16, 32))
        for ct_len in (PAGE_SZ - RESERVE_SZ - IV_SZ, PAGE_SZ - RESERVE_SZ):
            pt = bytes((i * 7 + 3) & 0xFF for i in range(ct_len))
            ct = AES.new(key, AES.MODE_CBC, iv).encrypt(pt)
            std = AES.new(key, AES.MODE_CBC, iv).decrypt(ct)
            raw = AES.new(key, AES.MODE_ECB).decrypt(ct)
            prev = int.from_bytes(iv + ct[:len(ct) - 16], "little")
            mine = (int.from_bytes(raw, "little") ^ prev).to_bytes(len(ct), "little")
            self.assertEqual(mine, std, f"ct_len={ct_len}")


# ── 数据安全修复：临时文件唯一性（D-2）─────────────────────────

class TestTempFileUniqueness(unittest.TestCase):
    """`_read_page1` / `decrypt_database` 的临时文件必须并发唯一。

    修复前用 f"siwx_p1_{os.getpid()}.tmp"：Flask 以 threaded=True 运行，
    `/api/status` 每次请求都会走 `collect_db_files` -> `_read_page1`，
    「微信占用中」时同进程多线程会撞名互相覆盖，读到对方的 page1，
    进而把 salt 张冠李戴写进密钥库。
    """

    def test_read_page1_concurrent_no_crosstalk(self):
        """并发读取多个文件，每个都必须拿到自己的 page1。"""
        import threading
        import builtins
        from siwx.sqlcipher import PAGE_SZ, _read_page1

        tmp = Path(tempfile.mkdtemp(prefix="siwx_p1_"))
        try:
            n = 8
            targets = []
            for i in range(n):
                p = tmp / f"db_{i}" / "message.db"
                p.parent.mkdir(parents=True, exist_ok=True)
                marker = 0x10 + i
                p.write_bytes(bytes([marker]) * 16
                              + bytes((j + marker) % 256 for j in range(PAGE_SZ - 16)))
                targets.append((p, marker))

            # 强制走「复制到临时文件」的回退分支
            real_open = builtins.open
            locked = {str(p) for p, _ in targets}

            def flaky_open(file, *a, **kw):
                if str(file) in locked:
                    raise OSError(13, "simulated lock")
                return real_open(file, *a, **kw)

            results, errors = {}, []
            lock = threading.Lock()
            barrier = threading.Barrier(n)

            def worker(path, marker):
                barrier.wait()
                try:
                    pg = _read_page1(path)
                    with lock:
                        results[str(path)] = pg[:16] if pg else None
                except Exception as e:      # noqa: BLE001
                    with lock:
                        errors.append(repr(e))

            builtins.open = flaky_open
            try:
                ts = [threading.Thread(target=worker, args=t) for t in targets]
                for t in ts:
                    t.start()
                for t in ts:
                    t.join()
            finally:
                builtins.open = real_open

            self.assertEqual(errors, [], f"并发异常: {errors}")
            for path, marker in targets:
                self.assertEqual(
                    results.get(str(path)), bytes([marker]) * 16,
                    f"{path.parent.name} 读到的 page1 不匹配（并发串扰）")

            residue = list(Path(tempfile.gettempdir()).glob("siwx_p1_*.tmp"))
            self.assertEqual(residue, [], f"临时文件残留: {residue}")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_temp_names_are_unique_across_calls(self):
        """mkstemp 生成的临时名必须唯一（旧实现恒为同一个名）。"""
        names = []
        for _ in range(20):
            fd, name = tempfile.mkstemp(prefix="siwx_p1_", suffix=".tmp")
            os.close(fd)
            names.append(name)
            Path(name).unlink(missing_ok=True)
        self.assertEqual(len(names), len(set(names)))

    def test_source_uses_mkstemp(self):
        """源码层面确认临时文件名不再是「固定前缀 + PID」的拼接。

        注意：注释里会引用旧实现的名字做说明，因此只看**代码行**
        （去掉注释与空行）是否还存在 `gettempdir() / f"siwx_..._{os.getpid()}"`。
        """
        src = Path(__file__).resolve().parent.parent / "siwx" / "sqlcipher.py"
        code_lines = []
        for line in src.read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if stripped.startswith("#") or not stripped:
                continue
            code_lines.append(stripped)
        code = "\n".join(code_lines)
        self.assertNotIn("siwx_p1_{os.getpid()}", code,
                         "仍在使用 PID 拼接的临时文件名")
        self.assertNotIn("siwx_db_{os.getpid()}", code,
                         "仍在使用 PID 拼接的临时文件名")
        self.assertIn("mkstemp", code)


# ── 数据安全修复：同名账号冲突检测（方案 C）────────────────────

class TestAccountConflicts(unittest.TestCase):

    def test_detects_duplicate_wxid(self):
        from siwx.discover import find_account_conflicts
        dirs = [
            ("wxid_a", r"C:\x\wxid_a\db_storage"),
            ("wxid_a", r"D:\x\wxid_a\db_storage"),
            ("wxid_b", r"C:\x\wxid_b\db_storage"),
        ]
        conflicts = find_account_conflicts(dirs)
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]["wxid"], "wxid_a")
        self.assertEqual(len(conflicts[0]["dirs"]), 2)

    def test_no_conflict_returns_empty(self):
        from siwx.discover import find_account_conflicts
        dirs = [("wxid_a", r"C:\x\a\db_storage"),
                ("wxid_b", r"C:\x\b\db_storage")]
        self.assertEqual(find_account_conflicts(dirs), [])

    def test_empty_input(self):
        from siwx.discover import find_account_conflicts
        self.assertEqual(find_account_conflicts([]), [])

    def test_three_copies(self):
        from siwx.discover import find_account_conflicts
        dirs = [("w", "1"), ("w", "2"), ("w", "3")]
        c = find_account_conflicts(dirs)
        self.assertEqual(len(c), 1)
        self.assertEqual(len(c[0]["dirs"]), 3)

    def test_status_api_exposes_conflicts(self):
        """/api/status 必须返回 conflicts 字段（前端据此提示）。"""
        from siwx import server
        client = server.app.test_client()
        resp = client.get("/api/status")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("conflicts", resp.get_json())


# ── 数据安全修复：manifest 来源保护（方案 B）──────────────────

class TestManifestSourceGuard(unittest.TestCase):
    """同名账号共用 output/<wxid>/ 时，来源变更不得静默覆盖已存在产物。

    核心兼容约束：旧 manifest 无 @source（升级自 v5.0.x）必须放行，
    行为与旧版完全一致。
    """

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="siwx_srcguard_"))
        from siwx import extract
        from siwx import pool
        self.extract, self.pool = extract, pool
        self._bench = []
        # 隔离密钥依赖：让 _resolve_key 必定成功，从而走到来源判定分支
        self._orig = (extract.parse_key, extract.verify_enc_key,
                      extract.decrypt_parallel)
        extract.parse_key = lambda k: b"\x00" * 32
        extract.verify_enc_key = lambda kb, p1: True

        def fake_parallel(tasks, workers=None, on_done=None):
            res = []
            for rel, src, dst, key_hex in tasks:
                Path(dst).parent.mkdir(parents=True, exist_ok=True)
                Path(dst).write_bytes(b"DECRYPTED")
                r = (rel, 2, "ok", "")
                res.append(r)
                if on_done:
                    on_done(r)
            return res

        extract.decrypt_parallel = fake_parallel
        self._bench.append(fake_parallel)

    def tearDown(self):
        (self.extract.parse_key, self.extract.verify_enc_key,
         self.extract.decrypt_parallel) = self._orig
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _entry(self, rel, path, size=8192):
        from siwx.sqlcipher import DbEntry
        return DbEntry(rel, Path(path), size, "aa" * 16, b"\x00" * 4096)

    def _make_dirs(self):
        c = self.tmp / "c" / "wxid_t" / "db_storage"
        d = self.tmp / "d" / "wxid_t" / "db_storage"
        for base in (c, d):
            (base / "contact").mkdir(parents=True, exist_ok=True)
            (base / "contact" / "contact.db").write_bytes(b"x" * 8192)
        return c, d

    def test_upgrade_without_source_field_is_allowed(self):
        """升级场景：旧 manifest 无 @source -> 放行（关键兼容性保证）。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)
        (out / "contact" / "contact.db").write_bytes(b"LEGACY")

        # 模拟 v5.0.2 的 manifest：只有业务键，没有 @source
        self.pool.save_manifest(out, {
            "contact\\contact.db": {"size": 1, "mtime": 1, "pages": 1,
                                    "key": "ab" * 32}})

        entries = [self._entry("contact\\contact.db", d / "contact" / "contact.db")]
        rep = self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                      entries=entries, use_cache=True)
        self.assertEqual(rep["conflicts"], 0, "升级用户不应被拦截")
        self.assertEqual(rep["ok"], 1)

    def test_source_change_blocks_overwrite(self):
        """来源变更 + 产物存在 -> 跳过，内容不变。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)
        target = out / "contact" / "contact.db"
        target.write_bytes(b"PROTECT-ME")

        self.pool.save_manifest(out, {
            self.extract.SOURCE_FIELD: str(c.resolve()).casefold()})

        entries = [self._entry("contact\\contact.db", d / "contact" / "contact.db")]
        rep = self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                      entries=entries, use_cache=True)
        self.assertEqual(rep["conflicts"], 1)
        self.assertEqual(target.read_bytes(), b"PROTECT-ME",
                         "其他副本的产物被覆盖了")

    def test_same_source_allows_decrypt(self):
        """来源一致 -> 正常解密。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)
        (out / "contact" / "contact.db").write_bytes(b"OLD")

        self.pool.save_manifest(out, {
            self.extract.SOURCE_FIELD: str(d.resolve()).casefold()})

        entries = [self._entry("contact\\contact.db", d / "contact" / "contact.db")]
        rep = self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                      entries=entries, use_cache=True)
        self.assertEqual(rep["conflicts"], 0)
        self.assertEqual(rep["ok"], 1)

    def test_writes_source_when_absent(self):
        """无历史来源且成功解密后，必须写入 @source。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)

        entries = [self._entry("contact\\contact.db", d / "contact" / "contact.db")]
        self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                entries=entries, use_cache=True)
        m = self.pool.load_manifest(out)
        self.assertEqual(m.get(self.extract.SOURCE_FIELD),
                         str(d.resolve()).casefold())

    def test_no_conflict_when_no_existing_artifact(self):
        """来源变更但无产物 -> 不拦截。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)   # 目录在，文件不在
        self.pool.save_manifest(out, {
            self.extract.SOURCE_FIELD: str(c.resolve()).casefold()})

        entries = [self._entry("contact\\contact.db", d / "contact" / "contact.db")]
        rep = self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                      entries=entries, use_cache=True)
        self.assertEqual(rep["conflicts"], 0)
        self.assertEqual(rep["ok"], 1)

    def test_source_field_coexists_with_business_keys(self):
        """@source 与业务键混存不互相干扰。"""
        out = self.tmp / "out" / "wxid_t"
        out.mkdir(parents=True)
        self.pool.save_manifest(out, {
            "contact\\contact.db": {"size": 1, "mtime": 1, "pages": 1,
                                    "key": "ab" * 32},
            self.extract.SOURCE_FIELD: "d:\\x"})
        m = self.pool.load_manifest(out)
        self.assertEqual(len(m), 2)
        self.assertIn("contact\\contact.db", m)
        self.assertEqual(m.get(self.extract.SOURCE_FIELD), "d:\\x")

    def test_report_includes_conflicts_count(self):
        """report 必须带 conflicts 字段（供上层/日志展示）。"""
        c, d = self._make_dirs()
        out = self.tmp / "out" / "wxid_t"
        (out / "contact").mkdir(parents=True, exist_ok=True)
        rep = self.extract.decrypt_dir(str(d), str(out), log=lambda m: None,
                                      entries=[], use_cache=True)
        self.assertIn("conflicts", rep)


class TestDisclaimerSync(unittest.TestCase):
    """免责声明双源同步：README 与控制台弹层（ui/pages/disclaimer.html）必须一致。

    免责条款改写时两处必须同步更新，避免「文档说一套、应用里另一套」。
    若调整哨兵条款的措辞，请同步修改本测试。
    """

    README = ROOT / "README.md"
    UI_DISCLAIMER = ROOT / "siwx" / "ui" / "pages" / "disclaimer.html"

    def test_ui_disclaimer_file_exists(self):
        self.assertTrue(self.UI_DISCLAIMER.is_file(),
                        "缺少 siwx/ui/pages/disclaimer.html（控制台免责弹层全文）")

    def test_key_clauses_present_in_both_sources(self):
        readme = self.README.read_text(encoding="utf-8")
        ui = self.UI_DISCLAIMER.read_text(encoding="utf-8")
        for phrase in (
            "技术研究与个人数据管理工具",
            "数据权属合法",
            "取得必要授权",
            "账号被平台限制或封禁",
            "世界多数国家和地区",
            "AS IS",
            "明文或不完全加密",
            "接入 AI 客户端前自行评估",
            "本声明不修改、不限制 AGPL-3.0 已授予的权利",
            "商用支持需另行授权",
            "可分割性与更新",
        ):
            self.assertIn(phrase, readme, f"README 免责声明缺少关键条款：{phrase}")
            self.assertIn(phrase, ui, f"应用内免责声明缺少关键条款：{phrase}")

    def test_consent_gate_wired_into_shell(self):
        app_js = (ROOT / "siwx" / "ui" / "app.js").read_text(encoding="utf-8")
        self.assertIn("DISCLAIMER_VERSION", app_js, "app.js 缺少免责条款版本常量")
        self.assertIn("siwx-disclaimer-ack", app_js, "app.js 未接入确认状态（localStorage）")
        self.assertIn("/pages/disclaimer.html", app_js, "app.js 未加载免责声明全文")
        index_html = (ROOT / "siwx" / "ui" / "index.html").read_text(encoding="utf-8")
        self.assertIn('id="side-disclaimer"', index_html, "侧栏缺少免责声明查看入口")


if __name__ == "__main__":
    unittest.main(verbosity=2)
